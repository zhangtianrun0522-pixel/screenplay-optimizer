import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import RGBColor
from docx.enum.text import WD_COLOR_INDEX


def set_run_font(run, is_chinese):
    run.font.name = "Times New Roman"
    if is_chinese:
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(ns + "rFonts")
        if rFonts is None:
            from lxml import etree
            rFonts = etree.SubElement(rPr, ns + "rFonts")
        rFonts.set(ns + "eastAsia", "宋体")


def add_colored_run(paragraph, text, color_rgb, strike=False):
    run = paragraph.add_run(text)
    is_chinese = any("一" <= char <= "鿿" for char in text)
    set_run_font(run, is_chinese)
    run.font.color.rgb = color_rgb
    if strike:
        run.font.strike = True


def add_yellow_highlight(paragraph, target_text):
    for run in paragraph.runs:
        if target_text in run.text:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def main():
    parser = argparse.ArgumentParser(description="生成标注版剧本 .docx")
    parser.add_argument("input_docx", type=str, help="原始剧本 .docx 路径")
    parser.add_argument("suggestions_json", type=str, help="建议列表 JSON 文件路径")
    args = parser.parse_args()

    doc = Document(args.input_docx)

    with open(args.suggestions_json, "r", encoding="utf-8") as f:
        suggestions = json.load(f)

    for sug in suggestions:
        p_idx = sug.get("paragraph_index", 0)
        if p_idx >= len(doc.paragraphs):
            continue
        para = doc.paragraphs[p_idx]
        s_type = sug.get("type", "")

        if s_type == "add":
            add_colored_run(para, sug["new_text"], RGBColor(0x00, 0x00, 0xFF))
        elif s_type == "delete":
            add_colored_run(para, sug["original_text"], RGBColor(0xFF, 0x00, 0x00), strike=True)
        elif s_type == "replace":
            add_colored_run(para, sug["original_text"], RGBColor(0xFF, 0x00, 0x00), strike=True)
            add_colored_run(para, sug["new_text"], RGBColor(0x00, 0x00, 0xFF))
        elif s_type == "comment":
            add_yellow_highlight(para, sug.get("original_text", ""))
            add_colored_run(para, f"【注：{sug['comment']}】", RGBColor(0xFF, 0xC0, 0x00))

    input_path = Path(args.input_docx)
    output_path = input_path.with_name(f"{input_path.stem}_annotated{input_path.suffix}")
    doc.save(str(output_path))
    print(f"已生成：{output_path}")


if __name__ == "__main__":
    main()
